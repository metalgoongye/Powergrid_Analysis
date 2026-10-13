# -*- coding: utf-8 -*-
"""귀속률 s — 지역소득 발간물 PDF에서 시도별 소득계정(수취 기준) 추출.

입력
  ① 2024지역소득통계.pdf p.226~ 「4. 소득계정 종합표 (2024p)」 (시도별 2쪽)
     - 짝수쪽: 비금융법인·금융법인·일반정부 + 표 제목(시도명)
     - 홀수쪽: 가계·총경제(합계) — 영문 행 라벨로 파싱
       행당 숫자 4개 = [가계 사용, 가계 원천, 총경제 사용, 총경제 원천]
  ② GRDP_GRI_지역외순수취본원소득 xlsx (param_s_attribution.py와 동일)

개념 주의 [확정 2026-10-13]
  이 표(본원소득분배계정)의 영업잉여·피용자보수는 **수취(제도부문 거주/등록지)
  기준**이다. 증거: 울산 영업잉여/GRDP 7.7% vs 서울 36.4% — 생산 기준이면
  불가능한 분포(중화학 울산이 최저일 수 없음). 울산 공장의 이윤이 서울 등록
  법인의 수취로 계상된 것으로, 이 자체가 이윤 역외 귀속의 직접 증거다.

검증(실패 시 멈춤)
  - 시도 17개 전부 발견
  - 시도별 총본원소득잔액(PDF, 십억원) == GRI(xlsx, 백만원) — 반올림 허용

산출
  시도별 수취 기준 계정 항목 CSV(분석 기초자료) + 구역 합산 요약.
  s의 직접 측정식: s ≈ 영업잉여수취₂ / 영업잉여발생₂
  → 분자(수취)는 여기서 확보. 분모(발생 기준 영업잉여)는 KOSIS
    '경제활동별 요소소득' 표(항목: 피용자보수/영업잉여)로 입수 후 결합한다.

출력: out/public/param_s_income_account_2024.csv + 콘솔. 단위: 조원 표시.
"""
import csv
import os
import re
import sys

from openpyxl import load_workbook
from pypdf import PdfReader

sys.stdout.reconfigure(encoding="utf-8")

DATA_ROOT = os.environ.get("DATA_ROOT", r"D:\z03.SmartGrid")
PDF = os.path.join(DATA_ROOT, "rdb", "GRDP_GRNI_2023_2024", "2024지역소득통계.pdf")
XLSX = os.path.join(DATA_ROOT, "rdb", "GRDP_GRNI_2023_2024",
                    "GRDP_GRI_지역외순수취본원소득_시도__20261008144507.xlsx")
OUT = os.path.join(os.path.dirname(__file__), "..", "out", "public",
                   "param_s_income_account_2024.csv")

SIDOS = ["서울특별시", "부산광역시", "대구광역시", "인천광역시", "광주광역시",
         "대전광역시", "울산광역시", "세종특별자치시", "경기도", "강원특별자치도",
         "충청북도", "충청남도", "전북특별자치도", "전라남도", "경상북도",
         "경상남도", "제주특별자치도"]
ZONE1 = {"서울특별시", "인천광역시", "경기도"}
PAGE_SCAN = range(215, 275)  # 종합표 수록 구간을 넉넉히 훑는다
TOL_MN = 200.0               # GRI 대조 허용오차(백만원) — PDF 0.1십억 반올림

ROWS = {  # 영문 행 라벨 -> 내부 키 (모두 수취/지급 기준, 총경제 열)
    "Compensation of employees": "comp",
    "Operating surplus": "os",
    "Net Taxes on production and imports": "ntax",
    "Property income": "prop",
    "Balance of primary incomes, gross": "gri",
}
NUM = re.compile(r"-?\d[\d ]*\.\d")


def parse_nums(prefix):
    """행 라벨 앞 토큰에서 숫자 4개([가계 사용, 가계 원천, 총경제 사용, 총경제 원천]).

    표의 빈칸은 단독 대시('-')로 찍히므로 0.0으로 치환 후 숫자를 뽑는다.
    (음수는 '-120.1'처럼 대시와 숫자가 붙어 나와 구분됨)
    """
    s = re.sub(r"(?<=\s)-(?=\s)", " 0.0 ", " " + prefix + " ")
    toks = NUM.findall(s)
    if len(toks) < 4:
        return None
    return [float(t.replace(" ", "")) for t in toks[-4:]]


def extract_pdf():
    r = PdfReader(PDF)
    found = {}      # 시도 -> {comp, os, os_hh, ntax, gri} (십억원)
    cur = None
    for i in PAGE_SCAN:
        txt = r.pages[i - 1].extract_text() or ""
        flat = re.sub(r"\s+", " ", txt)
        if "소득계정 종합표" in flat:
            cur = next((s for s in SIDOS if s in flat), None)
            continue
        if cur is None or cur in found:
            continue
        rec = {}
        for line in txt.splitlines():
            for label, key in ROWS.items():
                if label in line and key not in rec:
                    vals = parse_nums(line.split(label)[0])
                    if vals:
                        rec[key] = vals[3]          # 총경제 원천
                        if key == "os":
                            rec["os_hh"] = vals[1]  # 가계 원천(귀속임대 등)
                        if key == "prop":
                            rec["prop_use"] = vals[2]  # 총경제 지급
        if all(k in rec for k in ("comp", "os", "ntax", "prop", "gri")):
            found[cur] = rec
            cur = None
    return found


def main():
    # ① PDF 추출
    pdf = extract_pdf()
    missing = [s for s in SIDOS if s not in pdf]
    print(f"(1) PDF 추출: {len(pdf)}/17개 시도", "— 누락: " + ", ".join(missing) if missing else "")
    if missing:
        print("추출 실패 — 중단")
        return 1

    # ② xlsx 로드 + GRI 교차검증
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    rows = list(wb["데이터"].iter_rows(values_only=True))
    xl = {str(r[0]).strip(): (float(r[4]), float(r[7]), float(r[9]))
          for r in rows[2:] if r and r[0]}  # 시도 -> (GRDP24, GRI24, NR24) 백만원
    bad = 0
    for s in SIDOS:
        gri_pdf_mn = pdf[s]["gri"] * 1000.0  # 십억 -> 백만
        if abs(gri_pdf_mn - xl[s][1]) > TOL_MN:
            print(f"  GRI 불일치: {s} PDF {gri_pdf_mn:,.0f} vs xlsx {xl[s][1]:,.0f}")
            bad += 1
    if bad:
        print(f"(2) GRI 교차검증 실패 {bad}건 — 중단")
        return 1
    print(f"(2) GRI 교차검증: 17개 시도 전수 일치 (허용오차 {TOL_MN}백만원) — 통과")

    # ③ 시도별 수취 기준 계정 CSV (조원)
    rows_out = []
    for s in SIDOS:
        d = pdf[s]
        rows_out.append({
            "시도": s,
            "GRDP_조원": round(xl[s][0] / 1e6, 1),
            "피용자보수_수취_조원": round(d["comp"] / 1e3, 1),
            "영업잉여_수취_조원": round(d["os"] / 1e3, 1),
            "영업잉여_수취_가계분_조원": round(d["os_hh"] / 1e3, 1),
            "순생산수입세_조원": round(d["ntax"] / 1e3, 1),
            "재산소득_순수취_조원": round((d["prop"] - d["prop_use"]) / 1e3, 1),
            "GRI_조원": round(d["gri"] / 1e3, 1),
            "영업잉여수취_대_GRDP_pct": round(100 * d["os"] / 1e3 / (xl[s][0] / 1e6), 1),
        })

    # ④ 구역 합산 요약
    z2 = [s for s in SIDOS if s not in ZONE1]
    os2 = sum(pdf[s]["os"] for s in z2) / 1e3
    os2_hh = sum(pdf[s]["os_hh"] for s in z2) / 1e3
    grdp2 = sum(xl[s][0] for s in z2) / 1e6
    print(f"\n(3) 구역2(비수도권 14시도) 수취 기준 합산, 2024p (조원):")
    print(f"    영업잉여 수취 {os2:.1f} (가계분 {os2_hh:.1f}, 법인 등 {os2-os2_hh:.1f}) "
          f"| GRDP {grdp2:.1f}")
    print(f"    → s 직접 측정식: s ≈ 영업잉여수취2({os2:.1f}) / 영업잉여발생2(미입수)")
    print("    분모(발생 기준)는 KOSIS '경제활동별 요소소득' 표(피용자보수/영업잉여"
          " 항목 구분)로 입수 후 결합한다. 이 표의 수취 기준 수치로 s를 직접"
          " 계산하면 안 됨(개념 주의 참조).")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(rows_out[0].keys()))
        w.writeheader()
        w.writerows(rows_out)
    print(f"\n저장: {os.path.relpath(OUT)} (시도 17 × 수취 기준 계정 항목)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
