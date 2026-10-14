# -*- coding: utf-8 -*-
"""프로젝트 md 문서 → Word(docx) 일괄 변환.

제목·표·목록 구조를 유지하고 **굵게** 마크업을 굵은 런으로 바꾼다
(모델링 노트 학습판과 같은 꼴). 기본 출력은 docs/word/.

사용:
  python scripts/export_docs.py            # 전체
  python scripts/export_docs.py qa 명세    # 이름 일부로 선택
  python scripts/export_docs.py --dest "D:\\y06.워킹페이퍼\\..."  # 출력 폴더 지정
"""
import os
import re
import sys

from docx import Document
from docx.shared import Pt, RGBColor

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

TARGETS = {  # 키워드: (원본 md, 출력 파일명)
    "분석": ("docs/analysis_note.md", "분석노트.docx"),
    "qa": ("docs/qa_econ_concepts.md", "경제학개념QA.docx"),
    "명세": ("docs/data_provenance.md", "변수데이터명세.docx"),
    "geoai": ("docs/geoai_db_discussion_20261012.md", "GeoAI_DB논의_20261012.docx"),
}


def add_runs(p, text):
    for i, seg in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not seg:
            continue
        seg = re.sub(r"`([^`]*)`", r"\1", seg)  # 코드 마크업 제거
        r = p.add_run(seg)
        if i % 2 == 1:
            r.bold = True


def md_to_docx(src, dst):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "맑은 고딕"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
        "맑은 고딕")

    lines = open(src, encoding="utf-8").read().splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("# "):
            h = doc.add_heading(ln[2:].strip(), level=0)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0, 0, 0)
        elif ln.startswith("## "):
            doc.add_heading(ln[3:].strip(), level=1)
        elif ln.startswith("### "):
            doc.add_heading(ln[4:].strip(), level=2)
        elif ln.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].startswith("```"):
                buf.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(buf))
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        elif ln.startswith("|"):
            tbl = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    tbl.append(cells)
                i += 1
            i -= 1
            ncol = max(len(r_) for r_ in tbl)
            t = doc.add_table(rows=len(tbl), cols=ncol)
            t.style = "Table Grid"
            for ri, row in enumerate(tbl):
                for ci in range(ncol):
                    p = t.cell(ri, ci).paragraphs[0]
                    add_runs(p, row[ci] if ci < len(row) else "")
                    if ri == 0:
                        for r in p.runs:
                            r.bold = True
            doc.add_paragraph()
        elif re.match(r"^\s*- ", ln):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, re.sub(r"^\s*- ", "", ln))
        elif re.match(r"^\s*\d+\. ", ln):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\s*\d+\. ", "", ln))
        elif ln.strip() == "---":
            pass
        elif ln.strip():
            buf = [ln.strip()]
            while (i + 1 < len(lines) and lines[i + 1].strip()
                   and not re.match(r"^(#|\||\s*-\s|\s*\d+\. |---|```)", lines[i + 1])):
                i += 1
                buf.append(lines[i].strip())
            p = doc.add_paragraph()
            add_runs(p, " ".join(buf))
        i += 1
    doc.save(dst)


def main():
    args = [a for a in sys.argv[1:]]
    dest = os.path.join(ROOT, "docs", "word")
    if "--dest" in args:
        k = args.index("--dest")
        dest = args[k + 1]
        args = args[:k] + args[k + 2:]
    os.makedirs(dest, exist_ok=True)

    picks = {k: v for k, v in TARGETS.items()
             if not args or any(a.lower() in k for a in args)}
    if not picks:
        print("대상 없음. 키:", ", ".join(TARGETS))
        return 1
    for k, (src, name) in picks.items():
        full = os.path.join(ROOT, src)
        if not os.path.exists(full):
            print(f"건너뜀(원본 없음): {src}")
            continue
        out = os.path.join(dest, name)
        md_to_docx(full, out)
        print(f"저장: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
